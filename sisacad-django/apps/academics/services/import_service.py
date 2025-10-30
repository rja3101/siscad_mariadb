from django.contrib.auth.models import User, Group
from django.db import transaction
import openpyxl
from ..models import StudentProfile, Enrollment, Waitlist


class ImportService:
    """Servicio para importar datos desde XLSX"""
    
    @staticmethod
    def import_students(file):
        """
        Importa estudiantes desde archivo XLSX
        Formato: username, first_name, last_name, email
        """
        results = {
            'created': 0,
            'updated': 0,
            'errors': []
        }
        
        try:
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            
            # Obtener o crear grupo de estudiantes
            student_group, _ = Group.objects.get_or_create(name='Alumno')
            
            # Iterar filas (saltar encabezado)
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row[0]:  # Skip empty rows
                    continue
                
                try:
                    username = str(row[0]).strip()
                    first_name = str(row[1]).strip() if row[1] else ''
                    last_name = str(row[2]).strip() if row[2] else ''
                    email = str(row[3]).strip() if row[3] else f"{username}@example.com"
                    
                    with transaction.atomic():
                        # Crear o actualizar usuario
                        user, created = User.objects.get_or_create(
                            username=username,
                            defaults={
                                'first_name': first_name,
                                'last_name': last_name,
                                'email': email,
                                'is_active': True
                            }
                        )
                        
                        if not created:
                            # Actualizar datos
                            user.first_name = first_name
                            user.last_name = last_name
                            user.email = email
                            user.save()
                            results['updated'] += 1
                        else:
                            # Establecer password inicial
                            user.set_password('temporal123')
                            user.save()
                            results['created'] += 1
                        
                        # Agregar al grupo de estudiantes
                        user.groups.add(student_group)
                        
                        # Crear perfil si no existe
                        if not hasattr(user, 'student_profile'):
                            StudentProfile.objects.create(
                                user=user,
                                student_code=username,
                                cohort='2025-A',  # Default
                                curriculum='Plan 2020'  # Default
                            )
                
                except Exception as e:
                    results['errors'].append(f"Fila {row_num}: {str(e)}")
        
        except Exception as e:
            results['errors'].append(f"Error general: {str(e)}")
        
        return results
    
    @staticmethod
    def import_enrollments(file, term):
        """
        Importa matrículas desde archivo XLSX
        Formato: username, course_code, section
        """
        results = {
            'enrolled': 0,
            'waitlisted': 0,
            'errors': []
        }
        
        try:
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            
            from ..models import CourseGroup
            
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row[0]:  # Skip empty rows
                    continue
                
                try:
                    username = str(row[0]).strip()
                    course_code = str(row[1]).strip()
                    section = str(row[2]).strip()
                    
                    # Buscar usuario
                    try:
                        user = User.objects.get(username=username)
                    except User.DoesNotExist:
                        results['errors'].append(
                            f"Fila {row_num}: Usuario '{username}' no encontrado"
                        )
                        continue
                    
                    # Buscar grupo del curso
                    try:
                        course_group = CourseGroup.objects.get(
                            course__code=course_code,
                            section_code=section,
                            term=term
                        )
                    except CourseGroup.DoesNotExist:
                        results['errors'].append(
                            f"Fila {row_num}: Sección '{course_code}-{section}' no encontrada"
                        )
                        continue
                    
                    # Intentar matricular o agregar a waitlist
                    success, message = ImportService._enroll_or_waitlist(
                        user, course_group, term
                    )
                    
                    if 'waitlist' in message.lower():
                        results['waitlisted'] += 1
                    elif success:
                        results['enrolled'] += 1
                    else:
                        results['errors'].append(f"Fila {row_num}: {message}")
                
                except Exception as e:
                    results['errors'].append(f"Fila {row_num}: {str(e)}")
        
        except Exception as e:
            results['errors'].append(f"Error general: {str(e)}")
        
        return results
    
    @staticmethod
    @transaction.atomic
    def _enroll_or_waitlist(student, course_group, term):
        """Intenta matricular o agregar a waitlist"""
        
        # Verificar si ya está matriculado
        if Enrollment.objects.filter(
            student=student,
            course_group=course_group,
            term=term,
            status='ENROLLED'
        ).exists():
            return False, "Ya matriculado"
        
        # Verificar si ya está en waitlist
        if Waitlist.objects.filter(
            student=student,
            course_group=course_group,
            term=term,
            status='WAITING'
        ).exists():
            return False, "Ya en waitlist"
        
        # Verificar cupos
        from ..models import CourseGroup
        locked_group = CourseGroup.objects.select_for_update().get(pk=course_group.pk)
        
        if locked_group.get_available_slots() > 0:
            # Matricular
            Enrollment.objects.create(
                student=student,
                course_group=locked_group,
                term=term,
                status='ENROLLED'
            )
            return True, "Matriculado"
        else:
            # Agregar a waitlist
            last_position = Waitlist.objects.filter(
                course_group=course_group,
                term=term
            ).count()
            
            Waitlist.objects.create(
                student=student,
                course_group=course_group,
                term=term,
                position=last_position + 1,
                status='WAITING'
            )
            return True, "Agregado a waitlist"